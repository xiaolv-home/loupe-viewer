#!/usr/bin/env python3
"""Generates the sample files in fixtures/ used by tests and manual verification.

Needs: Pillow, fontTools (pip), ffmpeg (brew), plus macOS tools textutil, sips, bsdtar.
iWork / Office samples come from scripts/gen-iwork-fixtures.sh (needs Pages/Keynote/Numbers).
"""
import json, os, plistlib, random, shutil, struct, subprocess, sys, zlib
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent / "fixtures"
ROOT.mkdir(exist_ok=True)
random.seed(42)


def p(*parts):
    path = ROOT.joinpath(*parts)
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def run(*cmd):
    subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)


# ---------------------------------------------------------------- text & code
POEM = "床前明月光，疑是地上霜。\n举头望明月，低头思故乡。\n"
p("text", "中文-utf8.txt").write_text("UTF-8 中文文本\n" + POEM * 3, encoding="utf-8")
p("text", "中文-gbk.txt").write_bytes(("GBK 编码中文文本\n" + POEM * 3 + "繁體與简体混排：臺灣、台湾。\n").encode("gbk"))
p("text", "中文-gb18030.txt").write_bytes(("GB18030 扩展字符：𠀀𠀁 及常用字\n" + POEM).encode("gb18030"))
p("text", "big5.txt").write_bytes("繁體中文 Big5 編碼測試：臺北市、資訊、電腦。\n".encode("big5") * 5)
p("text", "utf16le-bom.txt").write_bytes("UTF-16 LE 带 BOM 的文本\n".encode("utf-16"))
p("text", "shift_jis.txt").write_bytes("日本語のテキスト、ひらがなとカタカナ。\n".encode("shift_jis") * 5)
p("text", "empty.txt").write_bytes(b"")
p("text", "crlf-windows.txt").write_bytes("Windows 换行\r\n第二行\r\n第三行\r\n".encode("utf-8"))

code = {
    "hello.py": '#!/usr/bin/env python3\n"""示例模块"""\nimport os\n\n@dataclass\nclass Point:\n    x: int = 0  # 横坐标\n\ndef main():\n    print(f"hello {os.getcwd()}")\n    return [i ** 2 for i in range(10)]\n',
    "app.js": "// 前端脚本\nconst greet = (name) => `你好, ${name}!`;\nexport default class App {\n  constructor() { this.count = 0x1F; }\n  async run() { return await fetch('/api'); }\n}\n",
    "types.ts": "interface User { id: number; name: string }\nexport function find(users: User[], id: number): User | undefined {\n  return users.find(u => u.id === id);\n}\n",
    "main.go": 'package main\n\nimport "fmt"\n\nfunc main() {\n\t// 注释\n\tfmt.Println("你好，世界")\n}\n',
    "lib.rs": "/// 文档注释\nfn main() {\n    let v: Vec<u32> = (0..10).collect();\n    println!(\"{:?}\", v);\n}\n",
    "Main.java": "public class Main {\n  public static void main(String[] args) {\n    System.out.println(\"中文\");\n  }\n}\n",
    "query.sql": "-- 查询门店\nSELECT id, name, COUNT(*) AS cnt\nFROM stores s\nLEFT JOIN orders o ON o.store_id = s.id\nWHERE s.state = 1\nGROUP BY id, name\nORDER BY cnt DESC LIMIT 10;\n",
    "deploy.sh": '#!/bin/bash\nset -euo pipefail\n# 部署脚本\nfor f in "$@"; do\n  echo "处理 ${f}"\ndone\n',
    "config.yaml": "# 配置\nserver:\n  host: 0.0.0.0\n  port: 8080\n  debug: true\nnames:\n  - 张三\n  - 李四\n",
    "settings.toml": '[package]\nname = "loupe"\nversion = "1.0.0"\n\n[deps]\nserde = { version = "1", features = ["derive"] }\n',
    "app.ini": "; INI 文件\n[database]\nhost = localhost\nport = 5432\n",
    "page.xml": '<?xml version="1.0" encoding="UTF-8"?>\n<root>\n  <!-- 注释 -->\n  <item id="1">中文</item>\n</root>\n',
    "style.css": "/* 样式 */\n.card { color: #ff6a3d; padding: 12px; }\n@media (prefers-color-scheme: dark) { body { background: #111; } }\n",
    "main.c": '#include <stdio.h>\nint main(void) {\n    printf("hi\\n");\n    return 0;\n}\n',
    "changes.diff": "--- a/file.txt\n+++ b/file.txt\n@@ -1,3 +1,3 @@\n 不变的行\n-删除的行\n+新增的行\n",
    "Makefile": "all: build\n\nbuild:\n\tswift build -c release\n",
    "Dockerfile": "FROM swift:6.0\nWORKDIR /app\nCOPY . .\nRUN swift build\n",
    "server.log": "".join(f"2026-10-22 10:{i // 60:02d}:{i % 60:02d}.123 {random.choice(['INFO', 'DEBUG', 'WARN', 'ERROR'])} [worker-{i % 4}] request {i} from 10.0.0.{i % 255} took {random.randint(1, 900)}ms\n" for i in range(400)),
}
for name, body in code.items():
    p("code", name).write_text(body, encoding="utf-8")
shutil.copy(Path(__file__).resolve().parent.parent / "Sources/Loupe/Renderers/TextRenderer.swift", p("code", "TextRenderer.swift"))

# Large file (> 50 MB) to exercise the virtualised line view.
big = p("text", "large-60mb.log")
if not big.exists() or big.stat().st_size < 50_000_000:
    with open(big, "w", encoding="utf-8") as f:
        line = 0
        while f.tell() < 62_000_000:
            line += 1
            f.write(f"{line:08d} 2026-10-22T08:00:00Z INFO 服务正常运行 payload={'x' * (line % 80)}\n")

# ---------------------------------------------------------------- markdown / html
p("docs", "README.md").write_text(r"""# Loupe 示例文档

这是一段 **Markdown**，包含 *斜体*、`行内代码` 与 [链接](https://example.com)。

## 表格

| 格式 | 渲染器 | 状态 |
|:-----|:------:|-----:|
| PDF | PDFKit | ✅ |
| 图片 | ImageIO | ✅ |

## 任务列表

- [x] 支持表格
- [x] 支持任务列表
- [ ] 支持魔法

## 公式

行内公式 $E = mc^2$，以及块级公式：

$$
\int_0^\infty e^{-x^2}\,dx = \frac{\sqrt{\pi}}{2}
$$

## 代码

```swift
struct Greeting {
    let text = "你好"
    func say() { print(text) }
}
```

## 图表

```mermaid
graph LR
  A[打开文件] --> B{识别格式}
  B -->|PDF| C[PDFKit]
  B -->|图片| D[ImageIO]
  B -->|未知| E[十六进制]
```

> 引用：文档中的脚本不会被执行。

<img src="https://example.com/remote.png" alt="远程图片（应被拦截）">
<script>document.body.innerHTML = "脚本被执行了！"</script>

![本地图片](../images/photo.jpg)
""", encoding="utf-8")

p("docs", "page.html").write_text("""<!doctype html>
<html lang="zh"><head><meta charset="utf-8"><title>HTML 示例</title>
<style>body{font-family:-apple-system;max-width:640px;margin:40px auto;line-height:1.6} .warn{color:#c00}</style></head>
<body><h1>HTML 页面示例</h1><p>这是一段 <b>HTML</b> 内容。</p>
<p id="js" class="warn">如果你看到这行字变成了“脚本已执行”，说明 JavaScript 没有被禁用。</p>
<script>document.getElementById('js').textContent = '脚本已执行';</script>
<img src="https://example.com/tracker.png" alt="远程资源应被拦截" width="1" height="1">
<ul><li>列表一</li><li>列表二</li></ul></body></html>
""", encoding="utf-8")

# ---------------------------------------------------------------- data
rows = [["门店", "城市", "日期", "营业额", "订单数", "备注"]]
cities = ["上海", "北京", "深圳", "杭州", "成都"]
for i in range(5000):
    rows.append([f"门店{i:04d}", random.choice(cities), f"2026-10-{i % 28 + 1:02d}", f"{random.uniform(1000, 99999):.2f}",
                 str(random.randint(10, 900)), "含逗号, 和\"引号\"" if i % 97 == 0 else ""])
def csv_line(r):
    out = []
    for c in r:
        if any(ch in c for ch in ',"\n'):
            c = '"' + c.replace('"', '""') + '"'
        out.append(c)
    return ",".join(out)
p("data", "sales-utf8.csv").write_text("\n".join(csv_line(r) for r in rows) + "\n", encoding="utf-8")
p("data", "sales-gbk.csv").write_bytes(("\r\n".join(csv_line(r) for r in rows[:300]) + "\r\n").encode("gbk"))
p("data", "table.tsv").write_text("\n".join("\t".join(r[:4]) for r in rows[:200]), encoding="utf-8")
big_csv = p("data", "large-200k.csv")
if not big_csv.exists():
    with open(big_csv, "w", encoding="utf-8") as f:
        f.write("id,name,value,ts\n")
        for i in range(200_000):
            f.write(f"{i},项目{i},{random.random():.6f},2026-10-22T00:00:{i % 60:02d}\n")

doc = {"name": "Loupe", "version": "1.0", "features": ["pdf", "image", "video"], "nested": {"a": 1, "b": [1, 2, {"c": None}], "中文键": "值"}, "enabled": True, "ratio": 0.75}
p("data", "config.json").write_text(json.dumps(doc, ensure_ascii=False, indent=2), encoding="utf-8")
p("data", "minified.json").write_text(json.dumps({"items": [{"id": i, "tags": ["a", "b"], "score": i * 1.5} for i in range(60)]}, separators=(",", ":")), encoding="utf-8")
p("data", "events.jsonl").write_text("\n".join(json.dumps({"event": e, "n": i}) for i, e in enumerate(["open", "close", "save"])), encoding="utf-8")
p("data", "broken.json").write_text('{"a": 1, "b": [1, 2,', encoding="utf-8")
plist_obj = {"CFBundleName": "Loupe", "Version": 3, "Enabled": True, "Items": ["一", "二"], "Data": b"\x00\x01\x02", "Nested": {"Key": "Value"}}
p("data", "binary.plist").write_bytes(plistlib.dumps(plist_obj, fmt=plistlib.FMT_BINARY))
p("data", "xml.plist").write_bytes(plistlib.dumps(plist_obj, fmt=plistlib.FMT_XML))

# ---------------------------------------------------------------- images
from PIL import Image, ImageDraw, ImageFont
def gradient(w, h):
    img = Image.new("RGB", (w, h))
    d = ImageDraw.Draw(img)
    for y in range(h):
        d.line([(0, y), (w, y)], fill=(int(255 * y / h), 90, 255 - int(255 * y / h)))
    for i in range(0, w, 40):
        d.line([(i, 0), (i, h)], fill=(255, 255, 255), width=1)
    return img

g = gradient(2400, 1600)
d = ImageDraw.Draw(g)
d.rectangle([100, 100, 500, 400], outline=(255, 255, 255), width=6)
g.save(p("images", "photo.jpg"), quality=92, exif=Image.Exif())
# JPEG with camera EXIF
exif = Image.Exif()
exif[0x010F] = "Leica"; exif[0x0110] = "Q3"; exif[0x0131] = "Loupe fixture"
ifd = exif.get_ifd(0x8769)
ifd[0x829A] = (1, 250); ifd[0x829D] = (17, 10); ifd[0x8827] = 400; ifd[0x920A] = (28, 1); ifd[0x9003] = "2026:10:22 09:30:00"
g.resize((1200, 800)).save(p("images", "exif-camera.jpg"), quality=90, exif=exif)
# Rotated by EXIF orientation 6 (should display upright portrait)
ex2 = Image.Exif(); ex2[0x0112] = 6
gradient(800, 400).save(p("images", "orientation-6.jpg"), exif=ex2)
rgba = Image.new("RGBA", (512, 512), (0, 0, 0, 0))
ImageDraw.Draw(rgba).ellipse([40, 40, 472, 472], fill=(255, 106, 61, 220))
rgba.save(p("images", "transparent.png"))
Image.new("RGB", (16, 16), (30, 30, 30)).save(p("images", "pixel-art.png"))
frames = [gradient(200, 200).rotate(a) for a in range(0, 360, 30)]
frames[0].save(p("images", "animated.gif"), save_all=True, append_images=frames[1:], duration=80, loop=0)
frames[0].save(p("images", "multipage.tiff"), save_all=True, append_images=frames[1:3])
g.resize((800, 533)).save(p("images", "photo.webp"), quality=85)
g.resize((800, 533)).save(p("images", "photo.bmp"))
rgba.resize((64, 64)).save(p("images", "icon.ico"))
run("sips", "-s", "format", "heic", str(p("images", "photo.jpg")), "--out", str(p("images", "photo.heic")))
try:
    g.resize((800, 533)).save(p("images", "photo.avif"))
except Exception:
    pass
p("images", "vector.svg").write_text("""<svg xmlns="http://www.w3.org/2000/svg" width="400" height="300" viewBox="0 0 400 300">
<rect width="400" height="300" fill="#f2efe9"/><circle cx="200" cy="150" r="90" fill="none" stroke="#d9481c" stroke-width="6"/>
<path d="M110 150H290M200 60V240" stroke="#222" stroke-width="2"/><text x="200" y="285" font-family="Helvetica" font-size="18" text-anchor="middle">SVG 矢量图</text>
<script>alert('svg script')</script></svg>""", encoding="utf-8")
# Huge image (> 16384 px side) to test downsampled display
Image.new("RGB", (20000, 400), (40, 120, 200)).save(p("images", "panorama-20000px.png"))

# ---------------------------------------------------------------- PDF with outline and text
def make_pdf(path, pages):
    objs = []
    def add(body):
        objs.append(body); return len(objs)
    font = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    page_ids, content_ids = [], []
    pages_id_placeholder = len(objs) + 2 * len(pages) + 1
    for i, text in enumerate(pages):
        stream = f"BT /F1 28 Tf 72 720 Td (Chapter {i + 1}) Tj ET BT /F1 14 Tf 72 680 Td ({text}) Tj ET"
        content_ids.append(add(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream"))
        page_ids.append(add(f"<< /Type /Page /Parent {pages_id_placeholder} 0 R /MediaBox [0 0 612 792] /Contents {content_ids[-1]} 0 R /Resources << /Font << /F1 {font} 0 R >> >> >>"))
    pages_id = add(f"<< /Type /Pages /Kids [{' '.join(f'{x} 0 R' for x in page_ids)}] /Count {len(page_ids)} >>")
    assert pages_id == pages_id_placeholder
    outline_id = len(objs) + 1
    item_ids = list(range(outline_id + 1, outline_id + 1 + len(page_ids)))
    add(f"<< /Type /Outlines /First {item_ids[0]} 0 R /Last {item_ids[-1]} 0 R /Count {len(item_ids)} >>")
    for i, pid in enumerate(page_ids):
        prev = f"/Prev {item_ids[i - 1]} 0 R" if i > 0 else ""
        nxt = f"/Next {item_ids[i + 1]} 0 R" if i < len(item_ids) - 1 else ""
        add(f"<< /Title (Chapter {i + 1}) /Parent {outline_id} 0 R {prev} {nxt} /Dest [{pid} 0 R /Fit] >>")
    catalog = add(f"<< /Type /Catalog /Pages {pages_id} 0 R /Outlines {outline_id} 0 R /PageMode /UseOutlines >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for o in offsets:
        out += f"{o:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objs) + 1} /Root {catalog} 0 R /Info << /Title (Loupe Sample) /Author (Loupe) >> >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(out)

make_pdf(p("docs", "outline.pdf"), [f"The quick brown fox jumps over the lazy dog. Page {i + 1} keyword loupe." for i in range(12)])
p("docs", "external-link.html")  # placeholder path ensures folder exists

# Office via textutil (rtf/doc/docx from HTML)
html_doc = p("docs", "_tmp.html")
html_doc.write_text("<meta charset=\"utf-8\"><h1>Word 文档示例</h1><p>这是由 textutil 生成的 <b>DOCX</b> 文档，包含中文。</p><ul><li>第一项</li><li>第二项</li></ul>", encoding="utf-8")
for fmt in ["docx", "rtf", "doc", "odt"]:
    run("textutil", "-inputencoding", "UTF-8", "-convert", fmt, str(html_doc), "-output", str(p("office", f"sample.{fmt}")))
html_doc.unlink()

# ---------------------------------------------------------------- archives
arch_src = ROOT / "_archive_src"
if arch_src.exists():
    shutil.rmtree(arch_src)
(arch_src / "项目" / "源码").mkdir(parents=True)
(arch_src / "项目" / "说明.txt").write_text("压缩包内的中文文件名", encoding="utf-8")
(arch_src / "项目" / "源码" / "main.swift").write_text('print("hi")\n')
(arch_src / "项目" / "data.bin").write_bytes(os.urandom(50_000))
for name, fmt in [("sample.zip", "zip"), ("sample.7z", "7zip"), ("sample.tar", "ustar"), ("sample.iso", "iso9660")]:
    run("bsdtar", "-c", "--format", fmt, "-f", str(p("archives", name)), "-C", str(arch_src), "项目")
run("bsdtar", "-czf", str(p("archives", "sample.tar.gz")), "-C", str(arch_src), "项目")
run("bsdtar", "-cJf", str(p("archives", "sample.tar.xz")), "-C", str(arch_src), "项目")
shutil.rmtree(arch_src)

# ---------------------------------------------------------------- fonts
for src, dst in [("/System/Library/Fonts/Supplemental/Arial.ttf", "Arial.ttf"),
                 ("/System/Library/Fonts/Supplemental/Songti.ttc", "Songti.ttc"),
                 ("/System/Library/Fonts/NewYork.ttf", "NewYork.ttf")]:
    if os.path.exists(src):
        shutil.copy(src, p("fonts", dst))
try:
    from fontTools.ttLib import TTFont
    f = TTFont("/System/Library/Fonts/Supplemental/Arial.ttf")
    f.flavor = "woff"
    f.save(p("fonts", "Arial.woff"))
except Exception as e:
    print("woff skipped:", e)

# ---------------------------------------------------------------- media (ffmpeg)
def ff(*args, out):
    if not p(*out).exists():
        subprocess.run(["ffmpeg", "-y", "-loglevel", "error", *args, str(p(*out))], check=False)

ff("-f", "lavfi", "-i", "testsrc2=size=1280x720:rate=30", "-f", "lavfi", "-i", "sine=frequency=440", "-t", "6", "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", out=("media", "h264-aac.mp4"))
ff("-f", "lavfi", "-i", "testsrc2=size=1280x720:rate=30", "-t", "4", "-c:v", "hevc_videotoolbox", "-tag:v", "hvc1", out=("media", "hevc.mov"))
ff("-f", "lavfi", "-i", "testsrc2=size=1280x720:rate=30", "-f", "lavfi", "-i", "sine=frequency=330", "-t", "6", "-c:v", "libx264", "-c:a", "ac3", out=("media", "h264-ac3.mkv"))
ff("-f", "lavfi", "-i", "testsrc2=size=640x360:rate=25", "-f", "lavfi", "-i", "sine=frequency=550", "-t", "5", "-c:v", "libvpx-vp9", "-b:v", "500k", "-c:a", "libopus", out=("media", "vp9-opus.webm"))
ff("-f", "lavfi", "-i", "testsrc2=size=640x360:rate=25", "-f", "lavfi", "-i", "sine=frequency=660", "-t", "4", "-c:v", "mpeg4", "-c:a", "mp3", out=("media", "mpeg4.avi"))
ff("-f", "lavfi", "-i", "testsrc2=size=640x360:rate=25", "-t", "4", "-c:v", "flv1", out=("media", "sample.flv"))
ff("-f", "lavfi", "-i", "sine=frequency=440:duration=5", "-c:a", "libmp3lame", "-metadata", "title=测试音频", "-metadata", "artist=Loupe", out=("media", "tone.mp3"))
ff("-f", "lavfi", "-i", "sine=frequency=523:duration=5", "-c:a", "flac", out=("media", "tone.flac"))
ff("-f", "lavfi", "-i", "sine=frequency=659:duration=3", out=("media", "tone.wav"))
ff("-f", "lavfi", "-i", "sine=frequency=784:duration=3", "-c:a", "libopus", out=("media", "tone.opus"))
p("media", "h264-aac.srt").write_text("1\n00:00:00,500 --> 00:00:02,500\n这是外挂 SRT 字幕\n\n2\n00:00:03,000 --> 00:00:05,500\n第二条字幕 <i>斜体</i>\n", encoding="utf-8")
p("media", "h264-aac.vtt").write_text("WEBVTT\n\n00:00.500 --> 00:02.500\nVTT 字幕\n", encoding="utf-8")
p("media", "h264-aac.ass").write_text("[Script Info]\nTitle: test\n\n[Events]\nFormat: Layer, Start, End, Style, Name, MarginL, MarginR, MarginV, Effect, Text\nDialogue: 0,0:00:01.00,0:00:04.00,Default,,0,0,0,,{\\b1}ASS 字幕{\\b0}\\N第二行\n", encoding="utf-8")

# ---------------------------------------------------------------- broken / mislabelled
p("edge", "truncated.png").write_bytes(p("images", "transparent.png").read_bytes()[:300])
p("edge", "garbage.pdf").write_bytes(b"%PDF-1.7\n" + os.urandom(4000))
p("edge", "not-a-zip.zip").write_bytes(b"PK\x03\x04" + os.urandom(2000))
p("edge", "truncated.mp4").write_bytes(p("media", "h264-aac.mp4").read_bytes()[:20000] if p("media", "h264-aac.mp4").exists() else b"\x00\x00\x00\x18ftypmp42")
shutil.copy(p("images", "transparent.png"), p("edge", "really-a-png.txt"))
shutil.copy(p("docs", "outline.pdf"), p("edge", "really-a-pdf.jpg"))
shutil.copy(p("data", "config.json"), p("edge", "no-extension"))
p("edge", "unknown.xyz").write_bytes(os.urandom(1024))
p("edge", "zero-bytes.bin").write_bytes(b"")
p("edge", "control-chars.txt").write_bytes(b"text with \x01\x02\x03 control bytes\x00 and NUL")
print("fixtures written to", ROOT)

# ---------------------------------------------------------------- spreadsheets (xlsx with formulas/dates, ods)
try:
    import datetime, zipfile
    from openpyxl import Workbook as XLWorkbook
    wb = XLWorkbook()
    ws = wb.active
    ws.title = "销售"
    ws.append(["门店", "城市", "日期", "营业额", "订单数", "客单价"])
    rnd = random.Random(7)
    for i in range(300):
        r = i + 2
        ws.append([f"门店{i:03d}", rnd.choice(["上海", "北京", "深圳", "杭州", "成都"]),
                   datetime.date(2026, 10, 1) + datetime.timedelta(days=i % 28),
                   round(rnd.uniform(1000, 50000), 2), rnd.randint(10, 500), f"=D{r}/E{r}"])
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.number_format = "yyyy-mm-dd"
    ws2 = wb.create_sheet("城市")
    ws2.append(["城市", "区域", "负责人"])
    for city, region, owner in [("上海", "华东", "张三"), ("北京", "华北", "李四"), ("深圳", "华南", "王五"), ("杭州", "华东", "赵六"), ("成都", "西南", "钱七")]:
        ws2.append([city, region, owner])
    wb.save(p("office", "sales.xlsx"))
except Exception as e:
    print("xlsx skipped:", e)

ods_content = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">
<office:body><office:spreadsheet>
<table:table table:name="库存">
<table:table-row><table:table-cell office:value-type="string"><text:p>商品</text:p></table:table-cell><table:table-cell office:value-type="string"><text:p>数量</text:p></table:table-cell><table:table-cell office:value-type="string"><text:p>单价</text:p></table:table-cell></table:table-row>
<table:table-row><table:table-cell office:value-type="string"><text:p>苹果</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="12"><text:p>12</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="3.5"><text:p>3.5</text:p></table:table-cell></table:table-row>
<table:table-row><table:table-cell office:value-type="string"><text:p>香蕉</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="30"><text:p>30</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="2"><text:p>2</text:p></table:table-cell></table:table-row>
<table:table-row table:number-rows-repeated="2"><table:table-cell office:value-type="string"><text:p>梨</text:p></table:table-cell><table:table-cell office:value-type="float" office:value="5"><text:p>5</text:p></table:table-cell><table:table-cell table:number-columns-repeated="3"/></table:table-row>
</table:table>
</office:spreadsheet></office:body></office:document-content>"""
with zipfile.ZipFile(p("office", "stock.ods"), "w") as z:
    z.writestr(zipfile.ZipInfo("mimetype"), "application/vnd.oasis.opendocument.spreadsheet")
    z.writestr("content.xml", ods_content)
    z.writestr("META-INF/manifest.xml", '<?xml version="1.0" encoding="UTF-8"?><manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/></manifest:manifest>')
print("spreadsheets written")

# ---------------------------------------------------------------- formatted report (styles, merges, widths)
try:
    import datetime
    from openpyxl import Workbook as XLWorkbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.styles.colors import Color
    wb = XLWorkbook()
    ws = wb.active
    ws.title = "月报"
    ws.merge_cells("A1:F1")
    ws["A1"] = "2026 年 10 月门店经营月报"
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="C0392B")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30
    headers = ["门店", "城市", "开业日期", "营业额", "同比", "备注"]
    thin = Side(style="thin", color="999999")
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=i, value=h)
        c.font = Font(bold=True, color="1F3864")
        c.fill = PatternFill("solid", fgColor=Color(theme=4, tint=0.6))   # accent1, lightened
        c.alignment = Alignment(horizontal="center")
        c.border = Border(bottom=thin)
    rows = [("南京西路店", "上海", datetime.date(2019, 3, 8), 1284567.5, 0.123, "旗舰店"),
            ("国贸店", "北京", datetime.date(2020, 7, 1), 986543.25, -0.052, ""),
            ("福田店", "深圳", datetime.date(2021, 11, 20), 765432.0, 0.287, "新装修"),
            ("西湖店", "杭州", datetime.date(2022, 5, 18), 432109.9, -0.181, "")]
    for r, row in enumerate(rows, 3):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
        ws.cell(row=r, column=3).number_format = 'yyyy"年"m"月"d"日"'
        ws.cell(row=r, column=4).number_format = '"¥"#,##0.00'
        ws.cell(row=r, column=5).number_format = '0.0%;[Red]-0.0%'
        if r % 2 == 0:
            for c in range(1, 7):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="F2F2F2")
    ws.cell(row=7, column=1, value="合计").font = Font(bold=True)
    ws.cell(row=7, column=4, value="=SUM(D3:D6)").number_format = '"¥"#,##0.00'
    ws.cell(row=7, column=4).font = Font(bold=True, italic=True, color="C0392B")
    for col, w in zip("ABCDEF", [16, 8, 20, 18, 10, 14]):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["F"].hidden = True
    ws.freeze_panes = "A3"
    wb.save(p("office", "report-styled.xlsx"))
except Exception as e:
    print("styled xlsx skipped:", e)
