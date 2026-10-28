#!/usr/bin/env python3
"""fixtures/office/structure.xlsx: formulas (incl. shared), merges, widths/heights,
conditional format, validation, defined name and a cross-sheet reference — for the
insert/delete rows & columns tests."""
import os, re, shutil, zipfile, tempfile
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Alignment, Font
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

out = os.path.join(os.path.dirname(__file__), "..", "fixtures", "office", "structure.xlsx")
wb = Workbook()
ws = wb.active
ws.title = "数据"
ws["A1"] = "季度销售"
ws.merge_cells("A1:D1")
ws["A1"].font = Font(bold=True, size=14)
ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 30
blue = PatternFill("solid", fgColor="4F81BD")
for i, h in enumerate(["名称", "数量", "单价", "金额"]):
    c = ws.cell(row=2, column=i + 1, value=h)
    c.fill = blue
    c.font = Font(bold=True, color="FFFFFF")
data = [("苹果", 3, 5.5), ("香蕉", 10, 2), ("橙子", 4, 6), ("葡萄", 2, 12)]
for r, (n, q, p) in enumerate(data, start=3):
    ws.cell(row=r, column=1, value=n)
    ws.cell(row=r, column=2, value=q)
    ws.cell(row=r, column=3, value=p)
    ws.cell(row=r, column=4, value=f"=B{r}*C{r}")
ws["A7"] = "合计"
ws["D7"] = "=SUM(D3:D6)"
ws["A8"] = "备注：数量单位为箱，单价单位为元，金额自动计算。"
ws.merge_cells("A8:B9")
ws["A8"].alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 16
ws.column_dimensions["D"].width = 12
ws.conditional_formatting.add("B3:B6", CellIsRule(operator="greaterThan", formula=["5"], fill=PatternFill("solid", fgColor="FFC7CE")))
dv = DataValidation(type="decimal", operator="greaterThan", formula1="0")
dv.add("C3:C6")
ws.add_data_validation(dv)
ws.freeze_panes = "A3"
s2 = wb.create_sheet("汇总")
s2["A1"] = "=数据!D7"
s2["A2"] = "=SUM(数据!B3:B6)"
s2["B1"] = "=数据!A3"
wb.defined_names["合计额"] = DefinedName("合计额", attr_text="数据!$D$7")
wb.save(out)

# openpyxl writes every formula in full; turn D3:D6 into one shared formula like Excel does.
tmp = tempfile.mkdtemp()
with zipfile.ZipFile(out) as z:
    z.extractall(tmp)
    names = z.namelist()
p = os.path.join(tmp, "xl", "worksheets", "sheet1.xml")
x = open(p, encoding="utf-8").read()
x = x.replace("<f>B3*C3</f>", '<f t="shared" ref="D3:D6" si="0">B3*C3</f>')
for r in (4, 5, 6):
    x = x.replace(f"<f>B{r}*C{r}</f>", '<f t="shared" si="0"/>')
open(p, "w", encoding="utf-8").write(x)
os.remove(out)
with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
    for n in names:
        z.write(os.path.join(tmp, n), n)
shutil.rmtree(tmp)
print("wrote", out)
